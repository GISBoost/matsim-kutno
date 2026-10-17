"""F3: popyt z GUS. Pracujący mieszkańcy per obwód (D5) i przyjeżdżający per gmina (D3, D9).

Wejścia (data/raw, sha256 w data/manifest.csv): tabela ludności BREC, tablice rynku pracy NSP 2021 (TABL. 5),
macierz przepływów, geometria BREC 2021 łódzkie.
Wyjścia: scenarios/kutno/demand/{workers_by_obwod,inbound_by_gmina}.csv, reports/f3_demand/f3_demand.{json,png}.
Uruchom: uv run python -W ignore pipeline/demand.py
"""
import json
import tomllib
from pathlib import Path

import geopandas as gpd
import matplotlib
import openpyxl
import pandas as pd
import pyogrio

matplotlib.use("Agg")
import matplotlib.pyplot as plt

cfg = tomllib.loads(Path("config/kutno.toml").read_text(encoding="utf-8"))["area"]
GMINA, POWIAT = cfg["teryt_gmina"], cfg["teryt_powiat"]
RAW = Path("data/raw")
OUT = Path("scenarios/kutno/demand")
OUT.mkdir(parents=True, exist_ok=True)

# --- 1. ludność per obwód (tabela BREC) ---------------------------------------------------------------
ws = [w for w in openpyxl.load_workbook(RAW / "gus/ludnosc_rejony_obwody_nsp2021.xlsx", read_only=True).worksheets
      if w.title.endswith("dzkie")][0]
rows = list(ws.iter_rows(values_only=True))
woj_row = next(r for r in rows if r[3] == "województwo")
i = next(k for k, r in enumerate(rows) if str(r[4]) == GMINA)
j = next(k for k in range(i + 1, len(rows)) if str(rows[k][3]).startswith("gmina") or rows[k][3] == "powiat")
# kolumny od 5: ogółem, M, K, przedprod., prod. razem, mobilny, niemobilny, poprod., 0-9 ... 80+
recs, rejon = [], None
for r in rows[i + 1:j]:
    if r[3] == "rejon statystyczny":
        rejon = int(r[4])
        continue
    assert r[3] == "obwód spisowy" and int(r[4]) < 10, r[:5]
    recs.append(dict(obwod_id=rejon * 10 + int(r[4]), rejon_id=rejon, pop=r[5], age_pre=r[8], age_work=r[9], age_post=r[12]))
ob = pd.DataFrame(recs)
assert (ob.age_pre + ob.age_work + ob.age_post == ob["pop"]).all()
gmina_pop = rows[i][5]

# --- 2. wskaźniki zatrudnienia: miasta woj. łódzkiego, NSP 2021 (TABL. 5) -----------------------------------
t5 = [[c for c in r if c is not None] for r in
      openpyxl.load_workbook(RAW / "gus/rynek_pracy_nsp2021_tablice.xlsx", read_only=True)["TABL. 5."].iter_rows(values_only=True)]
miasta = next(k for k, r in enumerate(t5) if r and str(r[0]).startswith("Miasta ogółem"))
blk = next(k for k in range(miasta, len(t5)) if t5[k] and str(t5[k][0]).strip() == "Łódzkie")
grp = {r[0].strip(): r for r in t5[blk:blk + 16] if r and isinstance(r[0], str)}
# r: [nazwa, ogółem, aktywni, pracujący, bezrobotni, bierni, nieustalony, ...]
def rates(name):
    r = grp[name]
    unknown = r[6] if isinstance(r[6], (int, float)) else 0
    return {"low": r[3] / r[1], "base": r[3] / (r[1] - unknown)}  # low: nieustalony = niepracujący; base: metoda GUS
rate = {"prod": rates("produkcyjny"), "poprod": rates("poprodukcyjny")}
# przedprodukcyjny w TABL. 5 to tylko 15-17 lat; w tabeli BREC grupa obejmuje 0-17. Wskaźnik odnosimy do 0-17 woj. łódzkiego
# (wiersz województwa w tabeli BREC, kolumna przedprodukcyjny), bo dla miast liczby 0-17 nie ma.
woj_t5 = next(r for r in t5 if r and str(r[0]).strip() == "Łódzkie")  # pierwszy blok = ogółem
woj_przed_15_17 = next(r for r in t5[t5.index(woj_t5):] if r and str(r[0]).strip() == "przedprodukcyjny")
rate["przed"] = {v: woj_przed_15_17[3] / woj_row[8] for v in ("low", "base")}

for v in ("low", "base"):
    ob[f"employed_{v}"] = ob.age_pre * rate["przed"][v] + ob.age_work * rate["prod"][v] + ob.age_post * rate["poprod"][v]

# --- 3. macierz dojazdów ---------------------------------------------------------------------------------
m = pd.read_excel(RAW / "gus/macierz_przeplywow_nsp2021.xlsx", sheet_name="Macierz przepływów", dtype={0: str, 2: str})
m.columns = ["home", "home_name", "work", "work_name", "n", "_"]
outgoing = int(m.loc[m.home == GMINA, "n"].sum())
inb = (m[m.work == GMINA].groupby(["home", "home_name"], as_index=False)["n"].sum().sort_values("n", ascending=False))
inb["in_powiat"] = inb.home.str.startswith(POWIAT)
inb["in_lodzkie"] = inb.home.str.startswith("10")
inb.to_csv(OUT / "inbound_by_gmina.csv", index=False)

# pracujący mieszkańcy, którzy pracują w Kutnie: odejmujemy wyjeżdżających proporcjonalnie we wszystkich obwodach
# (rozkład wyjeżdżających po obwodach nieznany). Macierz liczy tylko pracowników najemnych -> to dolna granica odpływu.
for v in ("low", "base"):
    tot = ob[f"employed_{v}"].sum()
    ob[f"work_in_kutno_{v}"] = ob[f"employed_{v}"] * (tot - outgoing) / tot
ob.round(2).to_csv(OUT / "workers_by_obwod.csv", index=False)

# --- 4. geometria + kontrola złączenia ------------------------------------------------------------------
g = pyogrio.read_dataframe(RAW / "brec/gml/AreaStatisticalUnit_10.xml",
                           columns=["LocationString", "thematicId|ThematicIdentifier|identifier"]).to_crs("EPSG:2180")
g["teryt"] = g["LocationString"].astype(str).str.zfill(7)
g["uid"] = g["thematicId|ThematicIdentifier|identifier"].astype(int)
geo_ob = g[(g.teryt == GMINA) & (g.uid > 999_999)]
geo_only = sorted(set(geo_ob.uid) - set(ob.obwod_id))
tab_only = sorted(set(ob.obwod_id) - set(geo_ob.uid))
assert not tab_only, tab_only

S = lambda c: round(float(ob[c].sum()))
report = {
    "population_gmina_row": gmina_pop, "population_sum_obwody": int(ob["pop"].sum()),
    "obwody_table": len(ob), "obwody_geometry": len(geo_ob), "obwody_geometry_only": geo_only,
    "join_rule": "obwod_id = rejon_id * 10 + nr_obwodu",
    "rates_source": "NSP 2021 Ludność na rynku pracy, TABL. 5, Miasta ogółem, Łódzkie (przed.: woj. 15-17 / 0-17)",
    "rates": {g_: {v: round(x, 4) for v, x in r.items()} for g_, r in rate.items()},
    "employed_residents": {"low": S("employed_low"), "base": S("employed_base")},
    "outgoing_employees": outgoing,
    "residents_working_in_kutno": {"low": S("work_in_kutno_low"), "base": S("work_in_kutno_base")},
    "inbound_total": int(inb.n.sum()), "inbound_gminy": len(inb),
    "inbound_in_powiat": int(inb[inb.in_powiat].n.sum()),
    "inbound_outside_powiat_in_lodzkie": int(inb[~inb.in_powiat & inb.in_lodzkie].n.sum()),
    "inbound_outside_lodzkie": int(inb[~inb.in_lodzkie].n.sum()),
    "agents_100pct": {v: S(f"work_in_kutno_{v}") + int(inb.n.sum()) for v in ("low", "base")},
}
report["ok"] = (report["population_sum_obwody"] == gmina_pop == 42255 and report["inbound_total"] == 8186
                and not tab_only)
json.dump(report, open("reports/f3_demand/f3_demand.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)

# --- 5. wizualizacja ---------------------------------------------------------------------------------
gob = geo_ob.merge(ob, left_on="uid", right_on="obwod_id", how="left")
gob["dens"] = gob.work_in_kutno_base / (gob.area / 1e6)
gm = g[g.teryt.str.startswith(POWIAT) & (g.uid <= 999_999)].dissolve(by="teryt").reset_index()
gm = gm.merge(inb[["home", "home_name", "n"]], left_on="teryt", right_on="home", how="left")

fig = plt.figure(figsize=(15, 7))
ax1 = fig.add_axes([0.02, 0.06, 0.33, 0.84])
gob.plot(ax=ax1, column="work_in_kutno_base", cmap="YlOrRd", edgecolor="0.5", lw=0.3, legend=True,
         legend_kwds={"shrink": 0.6, "label": "pracujący w Kutnie (bazowy) per obwód"},
         missing_kwds={"color": "0.85", "hatch": "///", "label": "bez danych"})
ax1.set_title(f"Kutno: {len(ob)} obwodów, pracujący mieszkańcy pracujący w Kutnie\n"
              f"bazowy {report['residents_working_in_kutno']['base']}, niski {report['residents_working_in_kutno']['low']}",
              fontsize=9)
ax1.set_axis_off()
ax2 = fig.add_axes([0.37, 0.06, 0.33, 0.84])
gm.plot(ax=ax2, column="n", cmap="Blues", edgecolor="0.3", lw=0.6, legend=True,
        legend_kwds={"shrink": 0.6, "label": "przyjeżdżający do pracy w Kutnie"})
for _, r in gm.iterrows():
    if pd.notna(r.n):
        p = r.geometry.representative_point()
        ax2.annotate(f"{r.home_name}\n{int(r.n)}", (p.x, p.y), ha="center", fontsize=6.5)
gpd.GeoSeries([geo_ob.union_all()]).boundary.plot(ax=ax2, color="#d1495b", lw=1)
ax2.set_title(f"przyjeżdżający z gmin powiatu: {report['inbound_in_powiat']} z {report['inbound_total']}", fontsize=9)
ax2.set_axis_off()
ax3 = fig.add_axes([0.80, 0.08, 0.18, 0.82])
top = inb[~inb.in_powiat].head(15)[::-1]
ax3.barh(top.home_name, top.n, color=["#2e86ab" if x else "#f18f01" for x in top.in_lodzkie])
ax3.set_title(f"spoza powiatu: {report['inbound_total'] - report['inbound_in_powiat']}\n"
              f"(pomarańczowe = spoza łódzkiego, razem {report['inbound_outside_lodzkie']})", fontsize=9)
ax3.tick_params(labelsize=7)
fig.suptitle(f"F3 popyt | agenci przy 100%: bazowy {report['agents_100pct']['base']}, niski {report['agents_100pct']['low']}"
             f" | {'ZALICZONA' if report['ok'] else 'NIE ZALICZONA'}", fontsize=11)
fig.savefig("reports/f3_demand/f3_demand.png", dpi=120)
print(json.dumps(report, indent=1, ensure_ascii=False))
