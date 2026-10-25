"""Krok 3 etapu 3 (D30): atrybuty osób (wiek, płeć, prawo jazdy, dostęp do auta) i pełna populacja Kutna.

Źródła (docs/ZRODLA.md, data/manifest.csv): BDL P4181 NSP 2021 (wiek pojedynczymi rocznikami × płeć: gm. Kutno, Łódzkie, Polska),
CEPiK 2019 (prawa jazdy wg płci i 6 grup wieku, Polska), GUS Mobilność 2024 (Tabl. 2, 3, 5: łódzkie), NSP 2021 rynek pracy TABL. 1
(wskaźnik zatrudnienia wg płci i wieku, miasta Polski), tabela ludności BREC (10-letnie grupy wieku w obwodach).
Metoda (docs/research/G_atrybuty_osob.md §6):
 1. udział praw jazdy r(płeć, grupa) = CEPiK / populacja Polski (mianownik: NSP 2021 = baza, Eurostat 1.01.2019 = wariant);
 2. przesunięcie w skali logit (jedna stała θ), by przy strukturze wieku Łódzkiego udział posiadaczy A lub B wśród 6+ = 60,6% (GUS);
 3. wiek × płeć pracujących = populacja (obwód: 10-letnie grupy BREC × rozkład gminy) × wskaźnik zatrudnienia (miasta Polski);
 4. dostęp do auta: bez prawa jazdy never; z prawem jazdy always z p = auta/posiadacze prawa jazdy (łódzkie, ZAŁOŻENIE: jedno auto = jeden
    posiadacz), inaczej sometimes.
Obecni agenci to jadący autem, więc dostają license=yes, carAvail=always albo sometimes (patrz Attrs.agents), a wiek z rozkładu ∝ populacja × zatrudnienie × P(licencja).
Użycie: uv run python -W ignore pipeline/attributes.py   (tworzy scenarios/kutno/demand/licence_rates.csv, population_kutno.csv
i reports/etap3/g_atrybuty/atrybuty.json)
"""
import json
import tomllib
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import brentq
from scipy.special import expit, logit

RAW = Path("data/raw")
G = RAW / "g_atrybuty"
GROUPS = [(18, 24), (25, 34), (35, 44), (45, 54), (55, 64), (65, 200)]  # grupy CEPiK
SEX = ("M", "K")
CAP = 0.98  # ZAŁOŻENIE: górne ograniczenie udziału z rejestru przed przesunięciem logit


def bdl():
    d = pd.read_csv(G / "bdl_P4181_nsp2021.csv", dtype={"age": str})
    d = d[(d.age != "ogółem") & (d.sex != "T")].copy()
    d["a"] = d.age.str.extract(r"(\d+)")[0].astype(int)  # „90 i więcej” -> 90
    return {u: g.pivot_table(index="a", columns="sex", values="pop", aggfunc="sum").reindex(index=range(91), columns=list(SEX)).fillna(0) for u, g in d.groupby("unit")}


def eurostat(year):
    d = pd.read_csv(G / "eurostat_demo_pjan_PL_2019_2021.csv")
    d = d[(d.time == year) & d.age.str.match(r"^Y\d+$|^Y_OPEN$")].copy()
    d["a"] = d.age.str.extract(r"(\d+)")[0].fillna(100).astype(int)
    d["sex"] = d.sex.map({"M": "M", "F": "K"})
    return d.pivot_table(index="a", columns="sex", values="pop", aggfunc="sum").reindex(columns=list(SEX)).fillna(0)


def group_pop(df, sex, lo, hi):
    return float(df.loc[lo:min(hi, df.index.max()), sex].sum())


def cepik():
    ws = openpyxl.load_workbook(G / "Prawa_jazdy_wg_stanu_na_31122019.xlsx", data_only=True).worksheets[0]
    out = {}
    for sex, age, n in list(ws.iter_rows(values_only=True))[1:]:
        s = "K" if str(sex).startswith("K") else "M"
        lo = 65 if str(age).startswith("OD") else int(str(age).split("-")[0])
        out[(s, [g[0] for g in GROUPS].index(lo))] = float(n)
    return out


def gus():
    wb = openpyxl.load_workbook(RAW / "gus_transport/mobilnosc_transportowa_ludnosci_w_polsce_w_2024_r._tablice.xlsx", read_only=True, data_only=True)
    row = lambda t, name: next(r for r in wb[t].iter_rows(values_only=True) if r[0] and str(r[0]).strip() == name)
    return {"lic_lodzkie": float(row("Tabl. 5", "Łódzkie")[1]), "members_lodzkie": float(row("Tabl. 3", "Łódzkie")[1]),
            "cars_lodzkie": float(row("Tabl. 2", "Łódzkie")[2]), "lic_pl": float(row("Tabl. 5", "Polska")[1]),
            "members_pl": float(row("Tabl. 3", "Polska")[1])}


def licence_rates(den="nsp2021", B=None, C=None, U=None):
    """r(płeć, grupa) po przesunięciu logit do 60,6% w Łódzkiem; zwraca (DataFrame, diagnostyka)."""
    B, C, U = B or bdl(), C or cepik(), U or gus()
    pl = B["Polska"] if den == "nsp2021" else eurostat(2019)
    raw = {k: C[k] / group_pop(pl, k[0], *GROUPS[k[1]]) for k in C}
    lod = B["Lodzkie"]
    target = U["lic_lodzkie"] / U["members_lodzkie"] * float(lod.loc[6:].to_numpy().sum())
    popl = {k: group_pop(lod, k[0], *GROUPS[k[1]]) for k in C}
    cap = {k: min(raw[k], CAP) for k in C}  # CEPiK liczy też zmarłych i wyjechanych: mężczyźni 65+ mają 1,04 > 1, więc rejestr jest górną granicą
    theta = brentq(lambda t: sum(expit(logit(cap[k]) + t) * popl[k] for k in C) - target, -5, 5)
    df = pd.DataFrame([dict(sex=k[0], group=f"{GROUPS[k[1]][0]}-{GROUPS[k[1]][1] if GROUPS[k[1]][1] < 200 else ''}", gi=k[1], cepik=C[k],
                            rate_raw=raw[k], rate=float(expit(logit(cap[k]) + theta))) for k in sorted(C, key=lambda k: (k[0], k[1]))])
    nat = B["Polska"]
    nat_share = sum(raw[k] * group_pop(nat, k[0], *GROUPS[k[1]]) for k in C) / float(nat.loc[6:].to_numpy().sum())
    return df, dict(den=den, theta=theta, target_lodzkie=target, raw_national_implied_share_6plus=nat_share,
                    gus_national_share=U["lic_pl"] / U["members_pl"], gus_lodzkie_share=U["lic_lodzkie"] / U["members_lodzkie"])


def lic_by_age(rates):
    """P(licencja | płeć, wiek 0..90): stała w grupie CEPiK, 0 poniżej 18 lat."""
    out = pd.DataFrame(0.0, index=range(91), columns=list(SEX))
    for r in rates.itertuples():
        lo, hi = GROUPS[r.gi]
        out.loc[lo:min(hi, 90), r.sex] = r.rate
    return out


def emp_rate():
    """Wskaźnik zatrudnienia (0-1) wg płci i wieku 0..90: miasta Polski, NSP 2021 TABL. 1 (rynek pracy)."""
    rows = list(openpyxl.load_workbook(RAW / "gus/rynek_pracy_nsp2021_tablice.xlsx", read_only=True, data_only=True)["TABL. 1"].iter_rows(values_only=True))
    i0 = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("Miasta"))
    out = pd.DataFrame(0.0, index=range(91), columns=list(SEX))
    for sex, label in (("M", "Mężczyźni"), ("K", "Kobiety")):
        i = next(i for i in range(i0, len(rows)) if rows[i][0] and str(rows[i][0]).startswith(label))
        grp = []
        for r in rows[i + 1:i + 20]:
            if r[0] and str(r[0]).strip()[0].isdigit():
                grp.append((int(str(r[0]).split()[0]), float(r[8]) / 100))
        for (lo, v), nxt in zip(grp, [g[0] for g in grp[1:]] + [91]):
            out.loc[lo:nxt - 1, sex] = v
    return out


def obwod_decades():
    ws = [w for w in openpyxl.load_workbook(RAW / "gus/ludnosc_rejony_obwody_nsp2021.xlsx", read_only=True).worksheets if w.title.endswith("dzkie")][0]
    rows = list(ws.iter_rows(values_only=True))
    i = next(k for k, r in enumerate(rows) if str(r[4]) == "1002011")
    j = next(k for k in range(i + 1, len(rows)) if str(rows[k][3]).startswith("gmina") or rows[k][3] == "powiat")
    out, rejon = {}, None
    for r in rows[i + 1:j]:
        if r[3] == "rejon statystyczny":
            rejon = int(r[4])
        elif r[3] == "obwód spisowy":
            out[rejon * 10 + int(r[4])] = np.array([0.0 if x in (None, "-") else float(x) for x in r[13:22]])  # 0-9 … 80+
    return out


def weights(kut, dec=None):
    """Waga komórki (wiek, płeć) w gminie albo, gdy podane 10-letnie grupy obwodu, w obwodzie (skalowanie dekad)."""
    w = kut.astype(float)
    if dec is not None:
        gdec = np.array([kut.loc[10 * g:10 * g + 9 if g < 8 else 90].to_numpy().sum() for g in range(9)])
        for g in range(9):
            sl = slice(10 * g, 10 * g + 9 if g < 8 else 90)
            w.loc[sl] = kut.loc[sl].to_numpy() * (dec[g] / gdec[g])
    return w


def car_p(mode, U):
    return {"cars_per_licence": min(1.0, U["cars_lodzkie"] / U["lic_lodzkie"]), "all_licence": 1.0}[mode]


class Attrs:
    def __init__(self, P):
        self.P = P
        self.B, self.U = bdl(), gus()
        self.rates, self.diag = licence_rates(P.get("licence_denominator", "nsp2021"), self.B, None, self.U)
        self.lic, self.emp = lic_by_age(self.rates), emp_rate()
        self.carq = car_p(P.get("car_avail_mode", "cars_per_licence"), self.U)
        self.kut = self.B["Kutno"]
        self.dec = obwod_decades()
        self.rng = np.random.default_rng(P.get("attr_seed", 30))

    def _draw(self, w, n):
        """n osób (wiek, płeć) z macierzy wag w (91 × 2)."""
        p = (w.to_numpy().ravel() / w.to_numpy().sum())
        k = self.rng.choice(p.size, size=n, p=p)
        return k // 2, np.array(SEX)[k % 2]

    def agents(self, meta):
        """Atrybuty obecnych agentów (jadący autem, license=yes): wiek i płeć ∝ populacja × zatrudnienie × P(licencja).
        Jadący autem to q = car_share × (1 − praca z domu) pracujących; wybierani najpierw spośród always (udział always wśród
        pracujących E_alw), reszta sometimes: P(always | agent) = min(1, E_alw / q)."""
        base = self.kut * self.emp
        q = self.P["car_share"] * (1 - self.P["work_from_home"])
        e_lic, e_all = float((base * self.lic).to_numpy().sum()), float(base.to_numpy().sum())
        self.p_alw = min(1.0, e_lic * self.carq / e_all / q)
        assert e_lic / e_all >= q, "pracujących z prawem jazdy mniej niż jadących autem: założenie car_share niespójne z udziałem prawa jazdy"
        age, sex = np.zeros(len(meta), int), np.empty(len(meta), object)
        for key, idx in meta.groupby(meta.kind.where(meta.kind != "resident", "r" + meta.key), sort=False).groups.items():
            dec = self.dec.get(int(str(key)[1:])) if str(key).startswith("r") and str(key)[1:].isdigit() else None
            a, s = self._draw(weights(self.kut, dec) * self.lic * self.emp, len(idx))
            pos = meta.index.get_indexer(idx)
            age[pos], sex[pos] = a, s
        car = np.where(self.rng.random(len(meta)) < self.p_alw, "always", "sometimes")
        return pd.DataFrame({"age": age, "sex": sex, "license": "yes", "carAvail": car}, index=meta.index)

    def population(self, workers):
        """Pełna populacja Kutna: dokładne 10-letnie sumy obwodów; pracujący wg `workers` (obwód -> liczba)."""
        rows = []
        for ob, dec in sorted(self.dec.items()):
            w = weights(self.kut, dec)
            for g in range(9):
                lo, hi = 10 * g, (10 * g + 9 if g < 8 else 90)
                ww = w.loc[lo:hi]
                n = int(round(dec[g]))
                if n == 0:
                    continue
                a, s = self._draw(ww.reindex(range(91)).fillna(0), n)
                rows += [(ob, int(x), y) for x, y in zip(a, s)]
        pop = pd.DataFrame(rows, columns=["obwod", "age", "sex"])
        pop["emp_w"] = [self.emp.at[a, s] for a, s in zip(pop.age, pop.sex)]
        pop["employed"] = False
        for ob, g in pop.groupby("obwod"):
            k = int(workers.get(ob, 0))
            if k:
                p = g.emp_w.to_numpy() / g.emp_w.sum()
                pop.loc[self.rng.choice(g.index.to_numpy(), size=k, replace=False, p=p), "employed"] = True
        lic = np.array([self.lic.at[a, s] for a, s in zip(pop.age, pop.sex)])
        pop["license"] = self.rng.random(len(pop)) < lic
        always = self.rng.random(len(pop)) < self.carq
        pop["carAvail"] = np.where(~pop.license, "never", np.where(always, "always", "sometimes"))
        return pop.drop(columns="emp_w")


def workers_scaled(total):
    ob = pd.read_csv("scenarios/kutno/demand/workers_by_obwod.csv")
    exact = ob.employed_base / ob.employed_base.sum() * total
    n = np.floor(exact).astype(int)
    n.iloc[np.argsort(-(exact - n).to_numpy(), kind="stable")[:int(round(exact.sum())) - int(n.sum())]] += 1
    return dict(zip(ob.obwod_id, n))


if __name__ == "__main__":
    cfg = tomllib.loads(Path("config/kutno.toml").read_text(encoding="utf-8"))
    P = cfg["plans"]
    Path("reports/etap3/g_atrybuty").mkdir(parents=True, exist_ok=True)
    rep = {"warianty": {}}
    for den in ("nsp2021", "eurostat2019"):
        for mode in ("cars_per_licence", "all_licence"):
            a = Attrs({**P, "licence_denominator": den, "car_avail_mode": mode})
            pop = a.population(workers_scaled(cfg["demand"]["employed_residents_reg"]))
            emp = pop[pop.employed]
            rep["warianty"][f"{den}/{mode}"] = dict(
                theta=round(a.diag["theta"], 4), p_always_given_licence=round(a.carq, 4),
                licence_share_6plus_kutno=round(float(pop[pop.age >= 6].license.mean()), 4), licence_share_18plus_kutno=round(float(pop[pop.age >= 18].license.mean()), 4),
                always_share_18plus_kutno=round(float((pop[pop.age >= 18].carAvail == "always").mean()), 4),
                employed=int(len(emp)), employed_license=round(float(emp.license.mean()), 4), employed_always=round(float((emp.carAvail == "always").mean()), 4),
                national_implied_share=round(a.diag["raw_national_implied_share_6plus"], 4), gus_national_share=round(a.diag["gus_national_share"], 4))
            if den == "nsp2021" and mode == "cars_per_licence":
                a.rates.to_csv("scenarios/kutno/demand/licence_rates.csv", index=False, lineterminator="\n")
                pop.insert(0, "person", [f"p{i}" for i in range(len(pop))])
                pop.to_csv("scenarios/kutno/demand/population_kutno.csv", index=False, lineterminator="\n")
                rep["baza"] = dict(diag=a.diag, rates=a.rates.round(4).to_dict("records"), n=len(pop))
                # piramida: populacja wg 10-letnich grup i płci vs BDL (TV)
                kut = a.kut
                dg = lambda df: np.array([df.loc[10 * g:(10 * g + 9 if g < 8 else 90)].to_numpy().sum() for g in range(9)])
                got = np.array([[((pop.age >= 10 * g) & (pop.age <= (10 * g + 9 if g < 8 else 90)) & (pop.sex == s)).sum() for g in range(9)] for s in SEX])
                rep["baza"]["pyramid_tv_pp"] = round(float(np.abs(got - np.array([dg(kut[s]) for s in SEX])).sum() / 2 / len(pop) * 100), 2)
                rep["baza"]["sex_male_share"] = round(float((pop.sex == "M").mean()), 4)
    Path("reports/etap3/g_atrybuty/atrybuty.json").write_text(json.dumps(rep, indent=1, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(rep, indent=1, ensure_ascii=False))
