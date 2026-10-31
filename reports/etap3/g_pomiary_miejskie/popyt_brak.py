"""Pomiar przed wdrożeniem (pkt 2 rekomendacji k1_zachod.md): ile pojazdów/dobę na K1–K5 dodałby popyt, którego R9.4 nie ma.

Grupa I: mieszkańcy Kutna, cele dom–X spoza modelu (odwiedziny, rekreacja, podwożenie, zdrowie).
Grupa II: mieszkańcy pozostałych gmin powiatu kutnowskiego (w obszarze sieci), wyjścia nie do pracy do Kutna:
          zakupy, usługi, zdrowie, odwiedziny, rekreacja (bez szkoły i podwożenia: krok 4B).
Liczba wyjść = stawka GUS Mobilność 2024 (tabl. 10, łódzkie: podróże dom–X / osoby 6+ z tabl. 4) × ludność 6+ × udział kierowcy auta.
Każde wyjście = 2 podróże (tam i z powrotem). Bez MATSim: trasy najszybsze w pustej sieci (Dijkstra, jak k1_zachod.py, kontrola K1 < 1%).
Liczymy przy udziale kierowcy = 1 i (II) udziale wyjść do Kutna p = 1, warianty to przeskalowanie liniowe (tabela w .md).

ZAŁOŻENIA (warianty w raporcie):
- udział kierowcy: 0,424 (errand_car_share D32) albo 0,479 (GUS tabl. 13, łódzkie: kierowca 1 894 401 / 3 958 088 podróży);
- II: p = udział wyjść nie do pracy mieszkańców gmin powiatu, które kończą się w Kutnie; brak danych, wariant 0,2 / 0,4 / 0,6;
- ludność 6+ = ogółem − 6/18 × wiek przedprodukcyjny (0–17), BREC NSP 2021; II tylko część gminy w obszarze sieci (wg obrysu budynków);
- domy: budynki BDOT jedno- i wielorodzinne, waga obrys × kondygnacje (jak D28, bez podziału na obwody);
- cele: wszystkie w gm. m. Kutno, losowe z wagą (jak D32), podwożenie do najbliższej szkoły podstawowej/przedszkola/żłobka (jak D35).
Użycie: uv run python -W ignore reports/etap3/g_pomiary_miejskie/popyt_brak.py   Wyjście: popyt_brak.{json,png} obok skryptu."""
import json
import tomllib
from pathlib import Path

import geopandas as gpd
import matplotlib
import numpy as np
import openpyxl
import pandas as pd
import pyogrio
import zstandard
from lxml import etree
from scipy.sparse import csr_matrix
from scipy.sparse.csgraph import dijkstra
from scipy.spatial import cKDTree

matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.style.use("dark_background")  # preferencja użytkownika

H = Path(__file__).parent
P = tomllib.load(open("config/kutno.toml", "rb"))["plans"]
rng = np.random.default_rng(2026)
z = lambda p: zstandard.ZstdDecompressor().stream_reader(open(p, "rb"))
CAR = {"0,424 (D32)": 0.424, "0,479 (GUS tabl. 13)": 1894401 / 3958088}
P_KUTNO = (0.2, 0.4, 0.6)

# --- stawki GUS (tabl. 10 i 4, łódzkie) -----------------------------------------------------------------------------
wb = openpyxl.load_workbook("data/raw/gus_transport/mobilnosc_transportowa_ludnosci_w_polsce_w_2024_r._tablice.xlsx", read_only=True)
t10 = next(r for r in wb["Tabl. 10"].iter_rows(values_only=True) if r[0] == "Łódzkie")
pop6_lodz = next(r for r in wb["Tabl. 4"].iter_rows(values_only=True) if r[0] == "Łódzkie")[1]
COL = {"zakupy_cc": 4, "zakupy_inne": 5, "uslugi": 6, "odwiedziny": 7, "rekreacja": 8, "podwozenie": 9, "zdrowie": 10}
RATE = {k: t10[c] / pop6_lodz for k, c in COL.items()}
GRUPA = {"I": ["odwiedziny", "rekreacja", "podwozenie", "zdrowie"],
         "II": ["zakupy_cc", "zakupy_inne", "uslugi", "zdrowie", "odwiedziny", "rekreacja"]}

# --- sieć i węzły przypięcia ----------------------------------------------------------------------------------------
nodes, L = {}, []
for _, el in etree.iterparse(z("output/r94c_gha/r94c-r93o-s4711/output_network.xml.zst"), tag=("node", "link")):
    if el.tag == "node":
        nodes[el.get("id")] = (float(el.get("x")), float(el.get("y")))
    elif "car" in el.get("modes"):
        a = {x.get("name"): x.text for x in el.iter("attribute")}
        L.append((el.get("id"), el.get("from"), el.get("to"), float(el.get("length")) / float(el.get("freespeed")) + 1, a.get("type")))
    el.clear()
L = pd.DataFrame(L, columns=["id", "fr", "to", "t", "type"]).set_index("id")
nid = {n: i for i, n in enumerate(nodes)}
XY = np.array(list(nodes.values()))
L["a"], L["b"] = L.fr.map(nid), L.to.map(nid)
best = L.reset_index().sort_values("t").drop_duplicates(["a", "b"])
M = csr_matrix((best.t, (best.a, best.b)), shape=(len(nid),) * 2)
ok = L[~L.type.fillna("").str.match(r"motorway|trunk|.*_link")]  # jak plans.py: bez A1, ekspresowych i łącznic
snap_nodes = np.unique(np.r_[ok.a, ok.b])
tree = cKDTree(XY[snap_nodes])
snap = lambda pts: snap_nodes[tree.query(np.c_[pts.x, pts.y])[1]]

# linki pomiarowe: K1–K5 (WIOŚ) i odcinki GPR (cel lekkie na kierunek z analiza/gpr_lcv.csv)
K = pd.read_csv("scenarios/kutno/validation/wios2013_linki.csv")
gp = pd.read_csv("output/r94c_gha/r94c-r93o-s4711/analiza/gpr_lcv.csv")
gp = gp[gp.link.isin(L.index)]
MEAS = {**{l: p for p, l in zip(K.pkt, K.link)}, **{l: f"GPR {n}" for n, l in zip(gp.nr, gp.link)}}
via_d = {l: dijkstra(M, indices=L.b[l]) for l in MEAS}  # z końca linku
via_o = {l: dijkstra(M.T.tocsr(), indices=L.a[l]) for l in MEAS}  # do początku linku

# --- budynki BDOT (cały powiat) --------------------------------------------------------------------------------------
bd = pyogrio.read_dataframe(P["home_bdot"], columns=["KATEGORIAISTNIENIA", "PRZEWAZAJACAFUNKCJABUDYNKU", "LICZBAKONDYGNACJI"])
bd = bd[bd.KATEGORIAISTNIENIA == "eksploatowany"].copy()
bd["pt"] = bd.geometry.representative_point()
bd["w"] = bd.geometry.area
bd = bd.set_geometry("pt")
gm = gpd.read_file("scenarios/kutno/area/gminy_powiat.gpkg").to_crs(2180).astype({"GMINA": int})
bd = gpd.sjoin(bd, gm[["GMINA", "geometry"]], how="inner", predicate="within")
area = gpd.GeoSeries.from_wkt([open("scenarios/kutno/area/network_area.wkt").read()], crs=2180).iloc[0]
bd["w_dom"] = bd.w * bd.LICZBAKONDYGNACJI.fillna(1).clip(lower=1)
dom = bd[bd.PRZEWAZAJACAFUNKCJABUDYNKU.isin(P["home_bdot_functions"])].copy()
dom["in_area"] = dom.within(area)
city = bd[bd.GMINA == 1002011]
F = lambda fs: city[city.PRZEWAZAJACAFUNKCJABUDYNKU.isin(fs)][["pt", "w"]].rename(columns={"pt": "geometry"}).set_geometry("geometry")
sh = pd.read_csv(P["errand_shops"], keep_default_na=False)
shop_pool = lambda c: gpd.GeoDataFrame({"w": 1.0}, index=sh.index[sh.shop.isin(["supermarket", "mall", "department_store"]) == c],
                                       geometry=gpd.points_from_xy(sh.x, sh.y)[sh.shop.isin(["supermarket", "mall", "department_store"]) == c])
POOL = {
    "zakupy_cc": shop_pool(True), "zakupy_inne": shop_pool(False),
    "uslugi": F(["bank", "inny urząd administracji publicznej", "urząd miasta", "urząd gminy", "starostwo powiatowe",
                 "placówka operatora pocztowego"]),  # jak errands.py TYPES["u"]
    "zdrowie": F(["placówka ochrony zdrowia", "szpital", "apteka"]),
    "odwiedziny": dom[dom.GMINA == 1002011][["pt", "w_dom"]].rename(columns={"pt": "geometry", "w_dom": "w"}).set_geometry("geometry"),
    "rekreacja": F(["dom kultury", "klub sportowy", "hala sportowa", "sala gimnastyczna", "basen kąpielowy", "teatr", "muzeum",
                    "biblioteka", "restauracja", "dom weselny"]),
    "podwozenie": F(["szkoła podstawowa", "przedszkole", "żłobek"]),
}

# --- ludność 6+ i wyjścia ---------------------------------------------------------------------------------------------
ws = [w for w in openpyxl.load_workbook("data/raw/gus/ludnosc_rejony_obwody_nsp2021.xlsx", read_only=True).worksheets if w.title.endswith("dzkie")][0]
brec = {int(r[4]): r[5] - r[8] * 6 / 18 for r in ws.iter_rows(values_only=True)
        if str(r[4]).startswith("1002") and r[3] in ("gmina miejska", "gmina wiejska", "miasto", "obszar wiejski")}
pop6_kutno = int((pd.read_csv("scenarios/kutno/demand/population_kutno.csv").age >= 6).sum())
origins = {"I": (dom[dom.GMINA == 1002011], {1002011: pop6_kutno})}
out = dom[(dom.GMINA != 1002011) & dom.in_area]
share_in = (out.groupby("GMINA").w_dom.sum() / dom[dom.GMINA != 1002011].groupby("GMINA").w_dom.sum()).to_dict()
origins["II"] = (out, {g: brec[g] * s for g, s in share_in.items()})


def tours(grp):
    """Wyjścia przy udziale kierowcy = 1 (i p = 1): DataFrame o (węzeł domu), d (węzeł celu), cel, gmina."""
    homes, pop = origins[grp]
    rows = []
    for cel in GRUPA[grp]:
        pool = POOL[cel]
        for g, n6 in pop.items():
            h = homes[homes.GMINA == g]
            n = rng.poisson(RATE[cel] * n6)
            o = h.iloc[rng.choice(len(h), n, p=(h.w_dom / h.w_dom.sum()).to_numpy())].pt
            if cel == "podwozenie":  # najbliższa placówka
                d = pool.geometry.iloc[cKDTree(np.c_[pool.geometry.x, pool.geometry.y]).query(np.c_[o.x, o.y])[1]]
            else:
                d = pool.geometry.iloc[rng.choice(len(pool), n, p=(pool.w / pool.w.sum()).to_numpy())]
            rows.append(pd.DataFrame({"o": snap(o), "d": snap(gpd.GeoSeries(d.to_numpy())), "cel": cel, "gmina": g}))
    return pd.concat(rows, ignore_index=True)


def count(tr):
    """Przejazdy przez linki pomiarowe (obie podróże wyjścia). Link na najszybszej trasie <=> t(o→a) + t_l + t(b→d) = t(o→d)."""
    res = pd.DataFrame(0.0, index=tr.index, columns=list(MEAS))
    for a_, b_ in (("o", "d"), ("d", "o")):
        for u, g in tr.groupby(a_):
            dist = dijkstra(M, indices=int(u))
            dd = dist[g[b_].to_numpy()]
            for l in MEAS:
                res.loc[g.index, l] += np.isclose(via_o[l][int(u)] + L.t[l] + via_d[l][g[b_].to_numpy()], dd, atol=1e-6)
    return res


R = {}
for grp in ("I", "II"):
    tr = tours(grp)
    tr = tr[tr.o != tr.d].reset_index(drop=True)
    c = count(tr)
    per = c.T.groupby(pd.Series(MEAS)).sum().T  # punkt (oba kierunki) na wyjście
    R[grp] = {"wyjscia_przy_kierowcy_1": tr.cel.value_counts().to_dict(), "punkty": per.sum().round(1).to_dict(),
              "punkty_wg_celu": per.groupby(tr.cel).sum().round(1).to_dict(orient="index")}
    if grp == "II":
        R[grp]["wyjscia_wg_gminy"] = tr.gmina.value_counts().to_dict()
        R[grp]["punkty_wg_gminy"] = per.groupby(tr.gmina).sum()[["K1", "K2", "K3", "K4", "K5"]].round(1).to_dict(orient="index")
    print(grp, len(tr), {k: round(v) for k, v in R[grp]["punkty"].items() if k.startswith("K")})

ref = json.load(open("reports/etap3/g_lancuchy/lancuchy_porownanie.json", encoding="utf-8"))["pomiary_miejskie"]
model, cel = ref["model_r93o_mediana"], ref["cel_2025"]
PK = ["K1", "K2", "K3", "K4", "K5"]
gpr_cel = gp.assign(k="GPR " + gp.nr.astype(str)).groupby("k").gpr_light_dir.sum()
war = {}
for cn, cs in CAR.items():
    i_ = {k: R["I"]["punkty"][k] * cs for k in PK}
    for p in P_KUTNO:
        ii = {k: R["II"]["punkty"][k] * cs * p for k in PK}
        tot = {k: model[k] + i_[k] + ii[k] for k in PK}
        gpr = {k: round(100 * (R["I"]["punkty"][k] + R["II"]["punkty"][k] * p) * cs / gpr_cel[k], 1) for k in gpr_cel.index}
        war[f"kierowca {cn}, p = {p}"] = {
            "I": {k: round(v) for k, v in i_.items()}, "II": {k: round(v) for k, v in ii.items()},
            "model_plus": {k: round(v) for k, v in tot.items()},
            "proc_celu": {k: round(100 * tot[k] / cel[k], 1) for k in PK},
            "suma_proc_celu": round(100 * sum(tot.values()) / sum(cel.values()), 1),
            "suma_K2_K5_proc_celu": round(100 * sum(tot[k] for k in PK[1:]) / sum(cel[k] for k in PK[1:]), 1),
            "gpr_dodatek_proc_celu_lekkie_mediana": float(np.median(list(gpr.values()))),
            "gpr_dodatek_proc_celu_lekkie_max": max(gpr.values()), "gpr_dodatek_proc_celu_lekkie": gpr}
out_j = {"stawki_gus_na_osobe6": {k: round(v, 5) for k, v in RATE.items()}, "pop6_kutno": pop6_kutno,
         "pop6_gmin_w_obszarze": {g: round(v) for g, v in origins["II"][1].items()}, "udzial_kierowcy": CAR,
         "model_r94": model, "cel_2025": cel, "przy_kierowcy_1_p_1": R, "warianty": war}
json.dump(out_j, open(H / "popyt_brak.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# --- wykres ----------------------------------------------------------------------------------------------------------
fig, ax = plt.subplots(1, 2, figsize=(16, 6), gridspec_kw={"width_ratios": [3, 2]})
cs = CAR["0,479 (GUS tabl. 13)"]
x = np.arange(len(PK))
ax[0].bar(x, [model[k] for k in PK], .6, color="#888", label="R9.4 (mediana 5 runów, GHA)")
b = np.array([model[k] for k in PK], float)
iv = np.array([R["I"]["punkty"][k] * cs for k in PK])
ax[0].bar(x, iv, .6, bottom=b, color="#4fc3f7", label="+ I: inne cele mieszkańców (kierowca 0,479)")
b += iv
for p0, p, col in zip((0,) + P_KUTNO[:-1], P_KUTNO, ("#a5d6a7", "#66bb6a", "#2e7d32")):
    seg = np.array([R["II"]["punkty"][k] * cs * (p - p0) for k in PK])  # narastająco od p0 do p
    ax[0].bar(x, seg, .6, bottom=b, color=col, label=f"+ II: przyjazdy z powiatu, p do {p}")
    b += seg
ax[0].scatter(x, [cel[k] for k in PK], color="w", marker="_", s=600, lw=2, label="cel 2025 (WIOŚ 2013 × 1,21)", zorder=3)
ax[0].scatter(x, [ref["cel_2025_korytarz"][k] for k in PK], color="w", marker="x", s=40, label="cel wg korytarza", zorder=3)
ax[0].set_xticks(x); ax[0].set_xticklabels(["K1 Toruńska", "K2 Łąkoszyńska", "K3 Grunwaldzka", "K4 Kościuszki", "K5 Łęczycka"])
ax[0].set_ylabel("poj./dobę, oba kierunki"); ax[0].set_title("pomiary miejskie: model + brakujący popyt (szacunek bez MATSim)")
ax[0].legend(fontsize=7)
g = pd.Series({k: (R["I"]["punkty"][k] + R["II"]["punkty"][k] * 0.4) * cs / gpr_cel[k] * 100 for k in gpr_cel.index}).sort_values()
ax[1].barh(g.index, g.to_numpy(), color="#ffb74d")
ax[1].axvline(15, color="w", ls=":", lw=1)
ax[1].set_xlabel("dodatek I + II (p = 0,4) w % celu GPR lekkie (suma kierunków)")
ax[1].set_title("ile z tego trafiłoby na odcinki GPR (tło po estymacji O-D by zmalało)")
ax[1].tick_params(axis="y", labelsize=7)
fig.suptitle("Brakujący popyt w mieście: inne cele mieszkańców (I) i przyjazdy nie do pracy z powiatu (II)")
fig.text(.01, .01, "Jak czytać. Lewy: słupek szary = natężenie R9.4 w punkcie WIOŚ (GHA run 41186041250); niebieski = dodatek grupy I "
         "(odwiedziny, rekreacja, podwożenie, zdrowie mieszkańców Kutna, GUS Mobilność 2024 tabl. 10 łódzkie × ludność 6+ × kierowca 0,479); "
         "zielone = dodatek grupy II (zakupy, usługi, zdrowie, odwiedziny, rekreacja mieszkańców gmin powiatu w obszarze sieci) przy udziale\n"
         "wyjść do Kutna p = 0,2 / 0,4 / 0,6 (ZAŁOŻENIE, segmenty narastająco). Biała kreska = cel 2025, krzyżyk = cel wg korytarza GPR. "
         "Prawy: ten sam dodatek (p = 0,4) na odcinkach GPR w % ich celu; kropki = 15%. Trasy najszybsze w pustej sieci (Dijkstra), bez MATSim "
         "i bez kongestii. Skrypt popyt_brak.py, opis w popyt_brak.md.", fontsize=8, va="bottom", wrap=True)
fig.tight_layout(rect=(0, .09, 1, 1)); fig.savefig(H / "popyt_brak.png", dpi=130)
